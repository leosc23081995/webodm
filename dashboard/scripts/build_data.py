#!/usr/bin/env python3
"""Build dashboard-data.js from the two source spreadsheets.

Reads:
  dashboard/data/WB2_Categoria_Completa.xlsx
  dashboard/data/PREVISAO_DE_CONSUMO_TUBOS.xlsx

Writes:
  dashboard/dashboard-data.js   (exposes window.DASH)

The shape preserves the keys consumed by dashboard-l0-l1.jsx, l2-l3.jsx,
l4.jsx and l5.jsx (intercompany, sazonalidade*, cap*, car*, months,
reconciliacao, despesas*, contratos, setores, anomalias, achados,
topClientesReun, kpis) and adds new keys with the rich content from the
spreadsheets (categorias, fabricantes, fabricantesMensal, equipamentos,
consumiveis, comodato, devolucao, transfInterna, previsao).
"""
from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
OUT_FILE = ROOT / "dashboard-data.js"

WB2 = DATA_DIR / "WB2_Categoria_Completa.xlsx"
PREV = DATA_DIR / "PREVISAO_DE_CONSUMO_TUBOS.xlsx"


def num(x, default=0.0):
    if x is None:
        return default
    try:
        return float(x)
    except (TypeError, ValueError):
        return default


def s(x):
    return None if x is None else str(x).strip()


def round2(x):
    return round(num(x), 2)


# ---------------------------------------------------------------------------
# 1. Load WB2 — categorized sheets
# ---------------------------------------------------------------------------
wb = openpyxl.load_workbook(WB2, data_only=True)

# Sheet 1: Summary by category × bank
sh_resumo = wb["1. Sumário Categorias"]
categorias = []
for row in sh_resumo.iter_rows(min_row=5, values_only=True):
    cat, banco, linhas, qtd, venda, sig = row[:6]
    if not cat:
        continue
    categorias.append({
        "categoria": s(cat),
        "banco": s(banco),
        "linhas": int(num(linhas)),
        "qtdItens": int(num(qtd)),
        "vendaTotal": round2(venda),
        "significado": s(sig),
    })

# Sheet 2: Manufacturers (537)
sh_fab = wb["2. Lista Fabricantes"]
fabricantes_all = []
for row in sh_fab.iter_rows(min_row=4, values_only=True):
    banco, cod, fab, qprod, qlin, qtot, venda = row[:7]
    if not fab:
        continue
    fabricantes_all.append({
        "banco": s(banco),
        "cod": s(cod),
        "fabricante": s(fab),
        "qtdProdutos": int(num(qprod)),
        "qtdLinhas": int(num(qlin)),
        "qtdTotal": int(num(qtot)),
        "venda29m": round2(venda),
    })
fabricantes_all.sort(key=lambda r: r["venda29m"], reverse=True)
fabricantes = fabricantes_all[:50]

# Sheet 3: Manufacturer × Category × Month
sh_fcm = wb["3. Fab x Mês x Categoria"]
fcm_rows = list(sh_fcm.iter_rows(values_only=True))
fcm_header = fcm_rows[2]
month_cols = [(idx, str(col)) for idx, col in enumerate(fcm_header) if col and len(str(col)) == 7 and str(col)[4] == "-"]
months = [m for _, m in month_cols]

fab_month_cat = []
for row in fcm_rows[3:]:
    fab, cat = row[0], row[1]
    if not fab:
        continue
    values = {m: round2(row[idx]) for idx, m in month_cols}
    fab_month_cat.append({
        "fabricante": s(fab),
        "categoria": s(cat),
        "values": values,
        "total": round2(row[-1]),
    })

# Aggregate per month per category (across all manufacturers) for intercompany / sazonalidade
def agg_by_month(filter_fn):
    out = {m: 0.0 for m in months}
    for r in fab_month_cat:
        if not filter_fn(r):
            continue
        for m in months:
            out[m] += r["values"].get(m, 0.0)
    return [{"month": m, "venda": round2(out[m])} for m in months]


# Sheet 4: Equipamentos
sh_eq = wb["4. EQUIPAMENTOS"]
eq_rows = list(sh_eq.iter_rows(values_only=True))
eq_header = eq_rows[2]
eq_month_cols = [(idx, str(col)) for idx, col in enumerate(eq_header) if col and len(str(col)) == 7 and str(col)[4] == "-"]
equipamentos_all = []
for row in eq_rows[3:]:
    cod, prod, fab = row[0], row[1], row[2]
    if not prod:
        continue
    values = {m: round2(row[idx]) for idx, m in eq_month_cols}
    equipamentos_all.append({
        "cod": s(cod),
        "produto": s(prod),
        "fabricante": s(fab),
        "values": values,
        "total": round2(row[-1]),
    })
equipamentos_all.sort(key=lambda r: r["total"], reverse=True)
equipamentos = equipamentos_all[:30]

# Sheet 5: Consumíveis top 500
sh_cons = wb["5. CONSUMIVEIS Top500"]
cons_rows = list(sh_cons.iter_rows(values_only=True))
cons_header = cons_rows[2]
cons_month_cols = [(idx, str(col)) for idx, col in enumerate(cons_header) if col and len(str(col)) == 7 and str(col)[4] == "-"]
consumiveis_all = []
for row in cons_rows[3:]:
    cod, prod, fab = row[0], row[1], row[2]
    if not prod:
        continue
    values = {m: round2(row[idx]) for idx, m in cons_month_cols}
    consumiveis_all.append({
        "cod": s(cod),
        "produto": s(prod),
        "fabricante": s(fab),
        "values": values,
        "total": round2(row[-1]),
    })
consumiveis_all.sort(key=lambda r: r["total"], reverse=True)
consumiveis = consumiveis_all[:50]

# Sheet 7: Comodato
sh_com = wb["7. COMODATO"]
com_rows = list(sh_com.iter_rows(values_only=True))
com_header = com_rows[2]
com_month_cols = [(idx, str(col)) for idx, col in enumerate(com_header) if col and len(str(col)) == 7 and str(col)[4] == "-"]
comodato_all = []
for row in com_rows[3:]:
    fab, prod = row[0], row[1]
    if not fab:
        continue
    values = {m: round2(row[idx]) for idx, m in com_month_cols}
    comodato_all.append({
        "fabricante": s(fab),
        "equipamento": s(prod),
        "values": values,
        "total": round2(row[-1]),
    })
comodato_all.sort(key=lambda r: r["total"], reverse=True)
comodato = comodato_all[:30]

# Sheet 8: Devolução
sh_dev = wb["8. DEVOLUCAO"]
dev_rows = list(sh_dev.iter_rows(values_only=True))
dev_header = dev_rows[2]
dev_month_cols = [(idx, str(col)) for idx, col in enumerate(dev_header) if col and len(str(col)) == 7 and str(col)[4] == "-"]
devolucao_all = []
for row in dev_rows[3:]:
    fab, prod = row[0], row[1]
    if not fab:
        continue
    values = {m: round2(row[idx]) for idx, m in dev_month_cols}
    devolucao_all.append({
        "fabricante": s(fab),
        "produto": s(prod),
        "values": values,
        "total": round2(row[-1]),
    })
devolucao_all.sort(key=lambda r: r["total"], reverse=True)
devolucao = devolucao_all[:30]

# Sheet 9: Transferência interna
sh_ti = wb["9. TRANSF INTERNA"]
ti_rows = list(sh_ti.iter_rows(values_only=True))
ti_header = ti_rows[2]
ti_month_cols = [(idx, str(col)) for idx, col in enumerate(ti_header) if col and len(str(col)) == 7 and str(col)[4] == "-"]
transf_all = []
for row in ti_rows[3:]:
    fab, prod = row[0], row[1]
    if not fab:
        continue
    values = {m: round2(row[idx]) for idx, m in ti_month_cols}
    transf_all.append({
        "fabricante": s(fab),
        "produto": s(prod),
        "values": values,
        "total": round2(row[-1]),
    })
transf_all.sort(key=lambda r: r["total"], reverse=True)
transf_interna = transf_all[:30]

# ---------------------------------------------------------------------------
# 2. Load PREVISAO — consumption forecast
# ---------------------------------------------------------------------------
wbp = openpyxl.load_workbook(PREV, data_only=True)
ws = wbp["RESUMO"]
# Row 5 has unit headers starting at column 8 (index 7 → 0-based)
header_row = list(ws.iter_rows(min_row=5, max_row=5, values_only=True))[0]
unit_headers = []
for idx in range(7, ws.max_column):
    name = header_row[idx]
    if name:
        unit_headers.append((idx, str(name).strip()))

previsao_produtos = []
total_estoque = 0
for row in ws.iter_rows(min_row=6, max_row=ws.max_row, values_only=True):
    desc = row[0]
    if not desc or str(desc).strip().lower().startswith("total"):
        continue
    estoque = int(num(row[1]))
    media_dia = num(row[2])
    media_mes = num(row[3])
    prev_dias = num(row[4])
    prev_meses = num(row[5])
    unidades = {}
    for idx, name in unit_headers:
        v = row[idx]
        if v is None or num(v) == 0:
            continue
        unidades[name] = int(num(v))
    total_estoque += estoque
    previsao_produtos.append({
        "produto": str(desc).strip(),
        "saldoEstoque": estoque,
        "mediaConsumoDia": round2(media_dia),
        "mediaConsumoMes": round2(media_mes),
        "previsaoDias": round2(prev_dias),
        "previsaoMeses": round2(prev_meses),
        "demandasUnidades": unidades,
    })

# MATRIZ sheet — supplier matrix per product
ws_mat = wbp["MATRIZ"]
matriz = []
for row in ws_mat.iter_rows(min_row=3, values_only=True):
    cod_supra, cod_fab, prod, fab, und, estoque = row[:6]
    if not prod:
        continue
    matriz.append({
        "codSupra": s(cod_supra),
        "codFab": s(cod_fab),
        "produto": s(prod),
        "fabricante": s(fab),
        "unidade": s(und),
        "estoque": int(num(estoque)),
    })

previsao = {
    "data": "14/05/2026",
    "periodo": "01/02/2026 a 30/04/2026",
    "totalEstoque": total_estoque,
    "produtos": previsao_produtos,
    "matriz": matriz,
}

# ---------------------------------------------------------------------------
# 3. Map to legacy DASH keys consumed by the JSX screens
# ---------------------------------------------------------------------------
def cat_filter(cat):
    return lambda r: r["categoria"] == cat

interco_monthly = agg_by_month(cat_filter("VENDA_INTERCO"))
venda_ext_monthly = agg_by_month(cat_filter("VENDA_EXTERNA"))
transf_monthly = agg_by_month(cat_filter("TRANSF_INTERNA"))
comodato_monthly = agg_by_month(cat_filter("COMODATO"))
devol_monthly = agg_by_month(cat_filter("DEVOLUCAO"))

# intercompany — month / venda; cost/margin unknown, leave 0
intercompany = [{
    "month": r["month"],
    "pedidos": 0,
    "itens": 0,
    "custo": 0,
    "venda": r["venda"],
    "margem": 0,
    "margemPct": 0,
} for r in interco_monthly]

# sazonalidade DESCART/REUNIDOS — split VENDA_EXTERNA per bank.
# We don't have bank breakdown by month, only by category total in sheet 1.
# Best effort: distribute proportional to category totals by bank.
ext_by_bank = {c["banco"]: c["vendaTotal"] for c in categorias if c["categoria"] == "VENDA_EXTERNA"}
total_ext = sum(ext_by_bank.values()) or 1.0
pct_descart = ext_by_bank.get("DESCART", 0) / total_ext
pct_reun = ext_by_bank.get("REUNIDOS", 0) / total_ext

sazonalidade_descart = [{
    "month": r["month"],
    "pedidos": 0,
    "itens": 0,
    "clientes": 0,
    "custo": 0,
    "venda": round2(r["venda"] * pct_descart),
    "margem": 0,
    "margemPct": 0,
} for r in venda_ext_monthly]

sazonalidade_reunidos = [{
    "month": r["month"],
    "pedidos": 0,
    "itens": 0,
    "clientes": 0,
    "custo": 0,
    "venda": round2(r["venda"] * pct_reun),
    "margem": 0,
    "margemPct": 0,
} for r in venda_ext_monthly]

# Contratos — derive from COMODATO top equipment (treat as active leasing)
contratos = []
for c in comodato[:30]:
    contratos.append({
        "empresa": "DESCART",
        "fornecedor": c["fabricante"],
        "setor": "—",
        "equipamento": c["equipamento"],
        "vencimento": "—",
        "diasRest": 0,
        "valorMensal": round2(c["total"] / max(1, sum(1 for v in c["values"].values() if v > 0))),
        "total29m": c["total"],
    })

# Top fabricantes (rebranded as topClientesReun lookalike)
top_fab = []
total_fab_venda = sum(f["venda29m"] for f in fabricantes) or 1.0
for f in fabricantes[:5]:
    top_fab.append({
        "nome": f["fabricante"],
        "origem": f["banco"],
        "total": f["venda29m"],
        "pct": round(f["venda29m"] / total_fab_venda, 4),
        "risco": "ALTO" if f["venda29m"] / total_fab_venda > 0.3 else "MEDIO",
    })

# KPIs derived
total_venda_externa = sum(c["vendaTotal"] for c in categorias if c["categoria"] == "VENDA_EXTERNA")
total_intercompany = sum(c["vendaTotal"] for c in categorias if c["categoria"] == "VENDA_INTERCO")
total_transf = sum(c["vendaTotal"] for c in categorias if c["categoria"] == "TRANSF_INTERNA")
total_comodato = sum(c["vendaTotal"] for c in categorias if c["categoria"] == "COMODATO")
total_devol = sum(c["vendaTotal"] for c in categorias if c["categoria"] == "DEVOLUCAO")
total_linhas = sum(c["linhas"] for c in categorias)
total_qtd = sum(c["qtdItens"] for c in categorias)

# Stub structure mirrors what L0–L5 access today (k.car, k.cap, k.despesa,
# k.resultado, intercompany.gap, contratos.{vigentes,mensalTotal,criticos},
# remuneracao.{gesta,claudio,leonardo,total}). Fields we cannot derive from
# the spreadsheets are kept at 0 — better explicit zero than a runtime crash.
contratos_mensal_total = sum(c.get("valorMensal", 0) for c in contratos)
kpis = {
    "reunidos": {
        "vendaExterna": round2(ext_by_bank.get("REUNIDOS", 0)),
        "car": round2(ext_by_bank.get("REUNIDOS", 0)),
        "cap": 0,
        "despesa": 0,
        "resultado": round2(ext_by_bank.get("REUNIDOS", 0)),
    },
    "descart": {
        "vendaExterna": round2(ext_by_bank.get("DESCART", 0)),
        "car": round2(ext_by_bank.get("DESCART", 0)),
        "cap": 0,
        "despesa": 0,
        "resultado": round2(ext_by_bank.get("DESCART", 0)),
    },
    "intercompany": {
        "valor": round2(total_intercompany),
        "gap": round2(total_intercompany),
    },
    "contratos": {
        "vigentes": len(contratos),
        "mensalTotal": round2(contratos_mensal_total),
        "criticos": sum(1 for c in contratos if c.get("valorMensal", 0) > 50000),
    },
    "remuneracao": {
        "gesta": 0,
        "claudio": 0,
        "leonardo": 0,
        "total": 0,
    },
    "transferenciaInterna": round2(total_transf),
    "comodato": round2(total_comodato),
    "devolucao": round2(total_devol),
    "meses": len(months),
    "totalRegistros": total_linhas,
    "totalItens": total_qtd,
    "totalFabricantes": len(fabricantes_all),
    "totalEquipamentos": len(equipamentos_all),
    "totalConsumiveis": len(consumiveis_all),
}

# Achados / anomalias — synthesized from the data
total_geral = sum(c["vendaTotal"] for c in categorias)
top_fab1 = fabricantes_all[0] if fabricantes_all else None
top_eq1 = equipamentos_all[0] if equipamentos_all else None
top_com1 = comodato_all[0] if comodato_all else None

achados = [
    {
        "num": "1",
        "titulo": "[CRÍTICO] TRANSF_INTERNA domina o volume financeiro",
        "evidencia": f"R$ {total_transf/1e6:.1f}M em transferências internas DESCART → unidades operacionais em 29 meses, contra R$ {total_venda_externa/1e6:.2f}M de venda externa.",
        "acao": "Validar se o regime de transferência está sendo precificado como custo / receita corretamente em DRE.",
    },
    {
        "num": "2",
        "titulo": "Concentração de fornecedor: " + (top_fab1["fabricante"] if top_fab1 else "—"),
        "evidencia": f"{top_fab1['fabricante']} = R$ {top_fab1['venda29m']/1e6:.2f}M ({top_fab1['banco']}) — {top_fab1['qtdProdutos']} produtos." if top_fab1 else "—",
        "acao": "Mapear contratos vigentes e SLAs do fornecedor principal; planejar segunda fonte.",
    },
    {
        "num": "3",
        "titulo": "Comodato relevante",
        "evidencia": f"COMODATO = R$ {total_comodato/1e6:.2f}M em equipamentos cedidos. Top: {top_com1['equipamento']} ({top_com1['fabricante']}) — R$ {top_com1['total']/1e6:.2f}M." if top_com1 else "—",
        "acao": "Revisar contrapartida em reagentes/consumíveis para cada equipamento em comodato.",
    },
    {
        "num": "4",
        "titulo": "Devolução acumulada",
        "evidencia": f"R$ {total_devol/1e6:.2f}M em devoluções em 29 meses ({len(devolucao_all)} SKUs).",
        "acao": "Análise causa-raiz das devoluções por SKU e por fabricante.",
    },
]

anomalias = []
# Equipment with single-month spike
for eq in equipamentos_all[:10]:
    vals = list(eq["values"].values())
    if not vals:
        continue
    mx = max(vals)
    if mx > 0 and mx / max(1.0, eq["total"]) > 0.85:
        # Spike: one month carries > 85% of total
        for m, v in eq["values"].items():
            if v == mx:
                anomalias.append({
                    "num": str(len(anomalias) + 1),
                    "titulo": f"Pico isolado: {eq['produto'][:60]}",
                    "evidencia": f"{m} concentrou R$ {v/1e6:.2f}M ({v/eq['total']*100:.0f}% do total 29m do SKU). Fabricante: {eq['fabricante']}.",
                    "acao": "Validar nota fiscal e categoria — possível lançamento único de comodato.",
                })
                break

# Products with low forecast in days (stock running out)
for p in previsao_produtos:
    if 0 < p["previsaoDias"] < 30 and p["mediaConsumoDia"] > 0:
        anomalias.append({
            "num": str(len(anomalias) + 1),
            "titulo": f"Estoque crítico: {p['produto']}",
            "evidencia": f"Saldo {p['saldoEstoque']} und cobre apenas {p['previsaoDias']:.0f} dias ao consumo médio de {p['mediaConsumoDia']:.0f}/dia.",
            "acao": "Disparar pedido de reposição — risco de ruptura em até 30 dias.",
        })

# Setores (não derivável → placeholder mínimo)
setores = []

# Reconciliação / CAP / CAR / Despesas — não cobertas pelas planilhas
reconciliacao = []
cap_descart = []
cap_reunidos = []
car_descart = []
car_reunidos = []
despesas_descart = []
despesas_reunidos = []

# ---------------------------------------------------------------------------
# 4. Compose DASH and emit
# ---------------------------------------------------------------------------
DASH = {
    # Legacy keys (mapped where possible)
    "months": months,
    "intercompany": intercompany,
    "sazonalidadeDescart": sazonalidade_descart,
    "sazonalidadeReunidos": sazonalidade_reunidos,
    "capDescart": cap_descart,
    "capReunidos": cap_reunidos,
    "carDescart": car_descart,
    "carReunidos": car_reunidos,
    "reconciliacao": reconciliacao,
    "despesasDescart": despesas_descart,
    "despesasReunidos": despesas_reunidos,
    "contratos": contratos,
    "setores": setores,
    "anomalias": anomalias,
    "achados": achados,
    "topClientesReun": top_fab,
    "kpis": kpis,
    # Novas keys derivadas das planilhas
    "categorias": categorias,
    "fabricantes": fabricantes,
    "fabricantesTodos": fabricantes_all,
    "fabricantesMensal": [
        {
            "fabricante": r["fabricante"],
            "categoria": r["categoria"],
            "values": r["values"],
            "total": r["total"],
        }
        for r in sorted(fab_month_cat, key=lambda x: x["total"], reverse=True)[:100]
    ],
    "equipamentos": equipamentos,
    "consumiveis": consumiveis,
    "comodato": comodato,
    "devolucao": devolucao,
    "transfInterna": transf_interna,
    "previsao": previsao,
    "_meta": {
        "geradoEm": "build_data.py",
        "fontes": [
            "dashboard/data/WB2_Categoria_Completa.xlsx",
            "dashboard/data/PREVISAO_DE_CONSUMO_TUBOS.xlsx",
        ],
        "periodo": f"{months[0]} a {months[-1]}" if months else "",
    },
}

OUT_FILE.write_text(
    "// Auto-generated by dashboard/scripts/build_data.py — do not edit by hand.\n"
    "window.DASH = " + json.dumps(DASH, ensure_ascii=False) + ";\n",
    encoding="utf-8",
)

print(f"Wrote {OUT_FILE} ({OUT_FILE.stat().st_size:,} bytes)")
print(f"  months: {len(months)} ({months[0] if months else '-'} → {months[-1] if months else '-'})")
print(f"  categorias: {len(categorias)}")
print(f"  fabricantes: {len(fabricantes_all)} (top {len(fabricantes)} kept)")
print(f"  equipamentos: {len(equipamentos_all)} (top {len(equipamentos)})")
print(f"  consumiveis: {len(consumiveis_all)} (top {len(consumiveis)})")
print(f"  comodato: {len(comodato_all)} (top {len(comodato)})")
print(f"  devolucao: {len(devolucao_all)} (top {len(devolucao)})")
print(f"  transfInterna: {len(transf_all)} (top {len(transf_interna)})")
print(f"  previsao produtos: {len(previsao_produtos)} / matriz: {len(matriz)}")
print(f"  achados: {len(achados)} / anomalias: {len(anomalias)}")
